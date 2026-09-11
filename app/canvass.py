"""
Canvass sheet (abstract of canvass) Excel builder.

Government/school procurement needs multiple supplier quotes per item;
this renders the N cheapest distinct-seller quotes per requested item into
a printable worksheet. Pure formatting — the price data comes from the
same fan-out search the /search endpoint uses.
"""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import NormalizedResult

_HEADERS = ["Item", "Rank", "Supplier", "Product Description", "UOM", "Price (PHP)", "Warranty", "Official", "URL"]
_WIDTHS = [28, 6, 24, 52, 10, 14, 20, 10, 60]

_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_CHEAPEST_FILL = PatternFill("solid", fgColor="E2EFDA")
_NO_QUOTE_FILL = PatternFill("solid", fgColor="FCE4EC")


def build_canvass_xlsx(item_quotes: list[tuple[str, list[NormalizedResult]]]) -> bytes:
    """item_quotes: (requested item, quotes sorted cheapest-first) per item."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Canvass"

    ws.append(["MARKET CANVASS SHEET"])
    ws.append([f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} — PRISM Price Aggregator"])
    ws.append([])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1).font = Font(italic=True, size=9, color="666666")

    header_row = ws.max_row + 1
    ws.append(_HEADERS)
    for col in range(1, len(_HEADERS) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center")

    for item, quotes in item_quotes:
        if not quotes:
            ws.append([item, "", "— walang nakuhang quote —", "", "", "", "", "", ""])
            for col in range(1, len(_HEADERS) + 1):
                cell = ws.cell(row=ws.max_row, column=col)
                cell.fill = _NO_QUOTE_FILL
                cell.border = _BORDER
            continue
        for rank, q in enumerate(quotes, start=1):
            ws.append([
                item if rank == 1 else "",
                rank,
                q.seller,
                q.description,
                q.uom,
                q.price_php,
                q.warranty,
                "YES" if q.is_official else "",
                q.url,
            ])
            row = ws.max_row
            for col in range(1, len(_HEADERS) + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = _BORDER
                if rank == 1:
                    cell.fill = _CHEAPEST_FILL
            ws.cell(row=row, column=1).font = Font(bold=True)
            ws.cell(row=row, column=6).number_format = "#,##0.00"

    for i, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
